"""
excel.py
========

Exportação dos dados coletados para Excel (.xlsx), CSV e JSON.

A planilha Excel é profissional:
  * cabeçalho colorido e em negrito;
  * primeira linha congelada;
  * autofiltro em todas as colunas;
  * largura de coluna ajustada automaticamente ao conteúdo;
  * colunas na ordem exata especificada em ``config.EXCEL_COLUMNS``.

As três exportações leem diretamente do banco SQLite, de modo que refletem o
estado atual da coleta a qualquer momento (inclusive durante uma pausa).
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import EXCEL_COLUMNS, Settings
from database import Database
from logger import get_logger
from models import OSC

log = get_logger("excel")

_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MAX_WIDTH = 60
_MIN_WIDTH = 10


def _row_from_record(record: dict[str, Any]) -> list[Any]:
    """Gera a linha de planilha a partir de um dict do banco (ordem EXCEL_COLUMNS)."""
    osc = OSC(**{k: record.get(k, "") for k in OSC.field_names() if k in record})
    # id_osc pode vir como int
    if record.get("id_osc") not in (None, ""):
        try:
            osc.id_osc = int(record["id_osc"])
        except (TypeError, ValueError):
            pass
    return osc.to_row(EXCEL_COLUMNS)


def export_excel(db: Database, path: Path) -> int:
    """Gera o arquivo .xlsx. Retorna o número de linhas exportadas."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws: Worksheet = wb.active
    ws.title = "OSCs"

    headers = [title for _key, title in EXCEL_COLUMNS]
    ws.append(headers)

    widths = [len(h) for h in headers]
    count = 0
    for record in db.iter_oscs():
        row = _row_from_record(record)
        ws.append(row)
        count += 1
        for i, value in enumerate(row):
            length = len(str(value)) if value is not None else 0
            if length > widths[i]:
                widths[i] = length

    # Estilo do cabeçalho.
    for col_idx, _title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
        cell.border = _BORDER

    # Larguras automáticas (limitadas).
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(
            _MIN_WIDTH, min(width + 2, _MAX_WIDTH)
        )

    # Congela cabeçalho e habilita autofiltro.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count + 1}"

    # Alinhamento do corpo.
    for row_cells in ws.iter_rows(min_row=2, max_row=count + 1, max_col=len(headers)):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    wb.save(path)
    log.info("Excel gerado: %s (%d linhas).", path, count)
    return count


def export_csv(db: Database, path: Path) -> int:
    """Gera o arquivo .csv (UTF-8 com BOM, para abrir no Excel PT-BR)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = [title for _key, title in EXCEL_COLUMNS]
    count = 0
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(headers)
        for record in db.iter_oscs():
            writer.writerow(_row_from_record(record))
            count += 1
    log.info("CSV gerado: %s (%d linhas).", path, count)
    return count


def export_json(db: Database, path: Path) -> int:
    """Gera o arquivo .json (lista de objetos com todos os campos canônicos)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    registros: list[dict[str, Any]] = []
    for record in db.iter_oscs():
        record.pop("raw_json", None)  # omite o payload bruto no JSON público
        registros.append(record)
    path.write_text(
        json.dumps(registros, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    log.info("JSON gerado: %s (%d registros).", path, len(registros))
    return len(registros)


def export_all(settings: Settings, db: Database | None = None) -> dict[str, int]:
    """Gera Excel, CSV e JSON de uma vez. Retorna a contagem por formato."""
    own = db is None
    db = db or Database(settings.db_path)
    try:
        return {
            "xlsx": export_excel(db, settings.xlsx_path),
            "csv": export_csv(db, settings.csv_path),
            "json": export_json(db, settings.json_path),
        }
    finally:
        if own:
            db.close()
