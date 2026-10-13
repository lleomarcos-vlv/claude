#!/usr/bin/env python3
"""Gera a apresentação da JP Alumínios: PDF, PDF leve, PPTX editável e JPGs por página."""
import hashlib
import html
import re
import shutil
import subprocess
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / "entregas"
BUILD = ROOT / "build"
CHROME = "/opt/pw-browsers/chromium"

CONTATO_NOME = "Felipe"
CONTATO_TEL = "(16) 99101-9797"
WA = "https://wa.me/5516991019797"
ARQ = "JP-Aluminios-Apresentacao"

INK = "#0E0F11"
INK2 = "#181A1D"
STEEL = "#22252A"
LINE = "#33373D"
LIGHT = "#F4F2EE"
MUTED = "#AEB2B8"
PAPER = "#F3F1EC"
LINE_L = "#D6D2C9"
MUTED_D = "#4E5258"
ACCENT = "#C9A26B"
ACCENT_D = "#8A6A3D"

W, H, M = 210, 297, 16
CW = W - 2 * M

FONTS = {
    ("display", True): ROOT / "fontes/BarlowCondensed-Bold.ttf",
    ("display", False): ROOT / "fontes/BarlowCondensed-Regular.ttf",
    ("body", True): ROOT / "fontes/Barlow-Bold.ttf",
    ("body", False): ROOT / "fontes/Barlow-Regular.ttf",
}
FAMILY = {"display": "Barlow Condensed", "body": "Barlow"}


# ---------------------------------------------------------------- elementos

def R(x, y, w, h, fill, alpha=1.0, radius=0, stroke=None, sw=0.3, link=None):
    return dict(k="rect", x=x, y=y, w=w, h=h, fill=fill, alpha=alpha, radius=radius,
                stroke=stroke, sw=sw, link=link)


def O(x, y, d, fill=None, stroke=None, sw=0.4):
    return dict(k="oval", x=x, y=y, w=d, h=d, fill=fill, stroke=stroke, sw=sw)


def I(x, y, w, h, src, focus=(0.5, 0.5), link=None):
    return dict(k="img", x=x, y=y, w=w, h=h, src=src, focus=focus, link=link)


def P(x, y, w, h, path, link=None):
    return dict(k="png", x=x, y=y, w=w, h=h, path=str(path), link=link)


def G(x, y, w, h, a0, a1, color=INK, ease=1.6):
    return dict(k="grad", x=x, y=y, w=w, h=h, a0=a0, a1=a1, color=color, ease=ease)


def T(x, y, w, h, text, size, color=LIGHT, font="body", bold=False, align="left",
      valign="top", lh=1.35, track=0.0, accent=ACCENT, pspace=0, link=None):
    return dict(k="text", x=x, y=y, w=w, h=h, text=text, size=size, color=color, font=font,
                bold=bold, align=align, valign=valign, lh=lh, track=track, accent=accent,
                pspace=pspace, link=link)


def A(x, y, w, h, url=WA):
    return dict(k="link", x=x, y=y, w=w, h=h, link=url)


def logo_w(h, dark=False):
    im = Image.open(ROOT / f"marca/logo_{'dark' if dark else 'white'}.png")
    return h * im.width / im.height


def eyebrow(y, text, color=ACCENT):
    return [R(M, y, 14, 0.7, color),
            T(M, y + 3, CW, 7, text, 13, color=color, font="display", bold=True, track=0.16)]


def footer(dark_text=True):
    c = LIGHT if dark_text else INK
    return [
        R(M, 278, CW, 0.3, LINE if dark_text else LINE_L),
        P(M, 283.4, logo_w(5.6), 5.6, ROOT / f"marca/logo_{'white' if dark_text else 'dark'}.png"),
        T(96, 281.2, W - M - 96, 10, f"Orçamentos com {CONTATO_NOME}  {{{CONTATO_TEL}}}", 13,
          color=c, font="display", bold=True, align="right", valign="middle", track=0.03,
          link=WA),
    ]


def tile_caption(x, y, w, h, title, sub=None):
    gh = h * 0.55
    els = [G(x, y + h - gh, w, gh, 0.0, 0.88)]
    if sub:
        els += [T(x + 5, y + h - 19.5, w - 10, 8.5, title, 17, font="display", bold=True, track=0.04),
                T(x + 5, y + h - 11.5, w - 10, 7, sub, 13.5, color="#D9D6D0")]
    else:
        els += [T(x + 4.5, y + h - 12, w - 9, 8.5, title, 15.5, font="display", bold=True,
                  track=0.05)]
    return els


# ---------------------------------------------------------------- páginas

def page_capa():
    e = [R(0, 0, W, H, INK),
         I(0, 0, W, 206, "012", focus=(0.5, 0.38)),
         G(0, 0, W, 64, 0.82, 0.0, ease=1.2),
         G(0, 96, W, 110, 0.0, 1.0, ease=1.5),
         P(M, 16, logo_w(10.5), 10.5, ROOT / "marca/logo_white.png"),
         T(W - M - 60, 16, 60, 10.5, "DESDE 2015", 13, color=ACCENT, font="display", bold=True,
           align="right", valign="middle", track=0.16)]
    e += eyebrow(176, "VIDROS E ESQUADRIAS DE ALUMÍNIO")
    e += [T(M, 187, CW, 40, "QUALIDADE QUE\nVALORIZA SEU {IMÓVEL.}", 56, font="display",
            bold=True, lh=0.94, track=0.005),
          T(M, 228, 168, 24, "Soluções sob medida para quem busca beleza, funcionalidade e um "
            "acabamento que faz a diferença.", 17, color="#D9D6D0", lh=1.38),
          R(M, 259, CW, 22, ACCENT, radius=1.6, link=WA),
          T(M + 7, 259, 90, 22, "SOLICITE SEU ORÇAMENTO\n{Fale com o " + CONTATO_NOME +
            " pelo WhatsApp}", 16, color=INK, font="display", bold=True, valign="middle",
            lh=1.15, track=0.06, accent="#3B2E1C", link=WA),
          T(M + 90, 259, CW - 97, 22, CONTATO_TEL, 30, color=INK, font="display", bold=True,
            align="right", valign="middle", lh=1.0, link=WA)]
    return dict(bg=INK, els=e)


def page_quem_somos():
    e = [R(0, 0, W, H, INK),
         I(0, 0, W, 122, "007", focus=(0.5, 0.55)),
         G(0, 62, W, 60, 0.0, 1.0, ease=1.4)]
    e += eyebrow(110, "QUEM SOMOS")
    e += [T(M, 121, CW, 30, "ESPECIALISTAS EM VIDROS\nE ESQUADRIAS DE {ALUMÍNIO.}", 42,
            font="display", bold=True, lh=0.96),
          T(M, 156, CW, 66,
            "A JP Alumínios nasceu com um objetivo claro: oferecer soluções em vidros e esquadrias "
            "de alumínio que unem durabilidade, segurança e elegância.\n"
            "Ao longo dos anos, construímos relações sólidas e de confiança com nossos clientes. "
            "É esse mesmo cuidado que queremos levar para o seu projeto.", 17, color="#D9D6D0", lh=1.42, pspace=9)]
    stats = [("2015", "ano de fundação"),
             ("3 SEGMENTOS", "residencial, comercial e industrial"),
             ("6 LINHAS", "de soluções em vidro e alumínio")]
    cw = (CW - 2 * 8) / 3
    for i, (t, d) in enumerate(stats):
        x = M + i * (cw + 8)
        e += [R(x, 231, cw, 0.6, ACCENT),
              T(x, 235, cw, 9, t, 23, font="display", bold=True, lh=1.0, track=0.02),
              T(x, 245, cw, 22, d, 14, color=MUTED, lh=1.32)]
    return dict(bg=INK, els=e + footer())


def page_servicos():
    e = [R(0, 0, W, H, INK)]
    e += eyebrow(16, "O QUE FAZEMOS")
    e += [T(M, 27, CW, 30, "SOLUÇÕES COMPLETAS\nEM VIDRO E {ALUMÍNIO}", 42, font="display",
            bold=True, lh=0.96),
          T(M, 59, CW, 17, "Do vão da janela à fachada inteira: conheça as linhas de produtos e "
            "serviços da JP Alumínios.", 16.5, color=MUTED, lh=1.38)]
    cards = [("ESQUADRIAS DE ALUMÍNIO", "Janelas, portas e venezianas sob medida para cada vão.",
              ("img", "013", (0.45, 0.5))),
             ("PELE DE VIDRO", "Fachadas envidraçadas, com mais luz natural e presença.",
              ("img", "050", (0.5, 0.45))),
             ("ESCADAS E GUARDA-CORPOS", "Corrimãos e guarda-corpos que protegem com visual leve.",
              ("img", "040", (0.5, 0.55))),
             ("BOX PARA BANHEIRO", "Box de vidro sob medida, prático e elegante no dia a dia.",
              ("img", "041", (0.5, 0.45))),
             ("ESPELHOS", "Espelhos sob medida que ampliam e valorizam os ambientes.",
              ("icon", "espelho", None)),
             ("COBERTURAS E PERGOLADOS", "Áreas externas protegidas e mais bem aproveitadas.",
              ("icon", "pergolado", None))]
    top, gap = 78, 5
    cw = (CW - 6) / 2
    ch = (274 - top - 2 * gap) / 3
    ph = ch - 27
    for i, (t, d, media) in enumerate(cards):
        x = M + (i % 2) * (cw + 6)
        y = top + (i // 2) * (ch + gap)
        e.append(R(x, y, cw, ch, INK2))
        if media[0] == "img":
            e.append(I(x, y, cw, ph, media[1], focus=media[2]))
        else:
            e += [R(x, y, cw, ph, STEEL), P(x + (cw - ph * 0.62) / 2, y + ph * 0.19, ph * 0.62,
                                            ph * 0.62, ROOT / f"build/icon_{media[1]}.png")]
        e += [T(x + 5, y + ph + 4, cw - 10, 7, t, 17, font="display", bold=True, track=0.03,
                lh=1.0),
              T(x + 5, y + ph + 11.5, cw - 10, 13, d, 13.5, color=MUTED, lh=1.3)]
    return dict(bg=INK, els=e + footer())


def page_detalhes():
    e = [R(0, 0, W, H, INK)]
    e += eyebrow(16, "QUALIDADE NOS DETALHES")
    e += [T(M, 27, CW, 30, "O ACABAMENTO\nFAZ A {DIFERENÇA.}", 42, font="display", bold=True,
            lh=0.96),
          T(M, 59, CW, 25, "Em vidro e alumínio, a qualidade aparece no alinhamento, no encaixe e "
            "no acabamento. Por isso, o cuidado começa na escolha dos materiais e vai até a "
            "instalação final.", 16.5, color=MUTED, lh=1.38)]
    top, bottom = 86, 272
    bw = 104
    e += [I(M, top, bw, bottom - top, "049", focus=(0.5, 0.45))]
    e += tile_caption(M, top, bw, bottom - top, "PORTAS DE ENTRADA",
                      "A primeira impressão de quem chega.")
    rx, rw = M + bw + 6, CW - bw - 6
    th = (bottom - top - 2 * 5) / 3
    tiles = [("037", (0.5, 0.5), "VENEZIANAS E BRISES"),
             ("044", (0.5, 0.42), "JANELAS SOB MEDIDA"),
             ("047", (0.5, 0.3), "GUARDA-CORPO DE VIDRO")]
    for i, (src, f, cap) in enumerate(tiles):
        y = top + i * (th + 5)
        e.append(I(rx, y, rw, th, src, focus=f))
        e += tile_caption(rx, y, rw, th, cap)
    return dict(bg=INK, els=e + footer())


def page_porque():
    e = [R(0, 0, W, H, PAPER)]
    e += eyebrow(16, "POR QUE A JP ALUMÍNIOS", color=ACCENT_D)
    e += [T(M, 27, CW, 30, "POR QUE CONSIDERAR\nA JP {ALUMÍNIOS?}", 42, color=INK,
            font="display", bold=True, lh=0.96, accent=ACCENT_D),
          T(M, 59, CW, 20, "Mais do que escolher uma janela ou um vidro, você está escolhendo quem "
            "vai cuidar de uma parte importante do seu imóvel.", 16.5, color=MUTED_D, lh=1.38)]
    items = [("MAIS DE 10 ANOS", "De experiência em vidros e esquadrias de alumínio, desde 2015."),
             ("EQUIPE ESPECIALIZADA", "Profissionais treinados para indicar a solução certa para "
              "cada projeto."),
             ("SOLUÇÕES SOB MEDIDA", "Personalizadas de acordo com as necessidades de cada "
              "cliente."),
             ("SEGURANÇA E CONFORTO", "Aliados à estética do seu projeto de arquitetura."),
             ("BOM CUSTO-BENEFÍCIO", "Orçamento otimizado, pensado para o que o seu projeto "
              "precisa."),
             ("ATENDIMENTO DIRETO", f"Você fala com o {CONTATO_NOME}, tira dúvidas e recebe "
              "orientação para escolher bem.")]
    cw = (CW - 10) / 2
    top, rh = 83, 46
    for i, (t, d) in enumerate(items):
        x = M + (i % 2) * (cw + 10)
        y = top + (i // 2) * rh
        e += [R(x, y, cw, 0.3, LINE_L),
              T(x, y + 3.5, 20, 10, f"{i + 1:02d}", 26, color=ACCENT_D, font="display",
                bold=True, lh=1.0),
              T(x, y + 14.5, cw, 7, t, 17.5, color=INK, font="display", bold=True, lh=1.0,
                track=0.03),
              T(x, y + 22.5, cw, 20, d, 14.5, color=MUTED_D, lh=1.32)]
    e += [I(0, 226, W, H - 226, "046", focus=(0.5, 0.13)),
          G(0, 226, W, H - 226, 0.62, 0.94, ease=1.0)]
    e += [T(M, 238, CW, 30, "Mais tranquilidade para decidir,\n{do orçamento à instalação.}", 24,
            font="display", bold=True, lh=1.05, track=0.01)]
    return dict(bg=PAPER, els=e + footer())


def page_aplicacoes():
    e = [R(0, 0, W, H, INK)]
    e += eyebrow(16, "APLICAÇÕES E POSSIBILIDADES")
    e += [T(M, 27, CW, 30, "IMAGINE ISSO NO\nSEU {PROJETO.}", 42, font="display", bold=True,
            lh=0.96),
          T(M, 59, CW, 18, "Construindo ou reformando? Cada ambiente pede uma solução diferente, e a "
            "JP Alumínios ajuda você a encontrar a sua.", 16.5,
            color=MUTED, lh=1.38)]
    tiles = [("045", (0.5, 0.5), "FACHADAS E ENTRADAS"),
             ("048", (0.35, 0.5), "LAZER E PISCINA"),
             ("015", (0.5, 0.5), "AMBIENTES INTERNOS"),
             ("023", (0.5, 0.5), "INTEGRAÇÃO COM A ÁREA EXTERNA")]
    top, bottom, gap = 84, 260, 6
    tw = (CW - gap) / 2
    th = (bottom - top - gap) / 2
    for i, (src, f, cap) in enumerate(tiles):
        x = M + (i % 2) * (tw + gap)
        y = top + (i // 2) * (th + gap)
        e.append(I(x, y, tw, th, src, focus=f))
        e += tile_caption(x, y, tw, th, cap)
    e += [T(M, 264, CW, 8, "CASAS  ·  OBRAS NOVAS  ·  REFORMAS  ·  COMÉRCIOS  ·  INDÚSTRIAS",
            13, color=ACCENT, font="display", bold=True, align="center", valign="middle",
            track=0.12)]
    return dict(bg=INK, els=e + footer())


def page_contato():
    e = [R(0, 0, W, H, INK),
         I(0, 0, W, 104, "038", focus=(0.5, 0.35)),
         G(0, 38, W, 66, 0.0, 1.0, ease=1.4)]
    e += eyebrow(95, "SOLICITE SEU ORÇAMENTO")
    e += [T(M, 106, CW, 32, "VAMOS TIRAR SEU\nPROJETO DO {PAPEL?}", 46, font="display",
            bold=True, lh=0.95),
          T(M, 140, CW, 16, f"Tudo pode começar com uma conversa. Conte para o "
            f"{CONTATO_NOME} o que você precisa e solicite seu orçamento.", 16.5,
            color="#D9D6D0", lh=1.38)]
    steps = [("CHAME NO WHATSAPP", f"Fale direto com o {CONTATO_NOME}."),
             ("CONTE SUA IDEIA", "Envie fotos, medidas ou o projeto, se tiver."),
             ("RECEBA SEU ORÇAMENTO", "Com orientação para a melhor solução.")]
    cw = (CW - 2 * 7) / 3
    for i, (t, d) in enumerate(steps):
        x = M + i * (cw + 7)
        y = 163
        e += [O(x, y, 10, stroke=ACCENT, sw=0.45),
              T(x, y, 10, 10, str(i + 1), 16, color=ACCENT, font="display", bold=True,
                align="center", valign="middle", lh=1.0),
              T(x, y + 13, cw, 7, t, 15.5, font="display", bold=True, lh=1.0, track=0.03),
              T(x, y + 20, cw, 13, d, 13.5, color=MUTED, lh=1.3)]
    py, ph = 202, 62
    qs = 46
    qx = W - M - 6 - qs
    e += [R(M, py, CW, ph, INK2, link=WA),
          R(M, py, 1.2, ph, ACCENT),
          T(M + 8, py + 7, 110, 7, f"{CONTATO_NOME.upper()}  |  JP ALUMÍNIOS", 13.5,
            color=ACCENT, font="display", bold=True, track=0.14),
          T(M + 8, py + 14, 116, 20, CONTATO_TEL, 46, font="display", bold=True, lh=1.0,
            valign="middle", link=WA),
          R(M + 8, py + 39, 82, 14, ACCENT, radius=7, link=WA),
          T(M + 8, py + 39, 82, 14, "CHAMAR NO WHATSAPP", 15, color=INK, font="display",
            bold=True, align="center", valign="middle", track=0.08, link=WA),
          R(qx, py + 5, qs, qs, "#FFFFFF", radius=1.2, link=WA),
          P(qx + 3.5, py + 8.5, qs - 7, qs - 7, ROOT / "build/qr_whatsapp.png", link=WA),
          T(qx, py + 52.5, qs, 6, "APONTE A CÂMERA", 11.5, color=MUTED, font="display",
            bold=True, align="center", valign="middle", track=0.14)]
    lw = logo_w(9)
    e += [P((W - lw) / 2, 272, lw, 9, ROOT / "marca/logo_white.png"),
          T(M, 284, CW, 6, "DESDE 2015  ·  RIBEIRÃO PRETO/SP", 12, color=MUTED, font="display",
            bold=True, align="center", valign="middle", track=0.16)]
    return dict(bg=INK, els=e)


PAGES = [page_capa, page_quem_somos, page_servicos, page_detalhes, page_porque,
         page_aplicacoes, page_contato]


# ---------------------------------------------------------------- recursos gerados

def hexrgb(c):
    c = c.lstrip("#")
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


def make_icons():
    col = hexrgb(ACCENT) + (255,)
    s = 1600
    sw = 34
    im = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    x0, x1, y0, y1 = 470, 1130, 170, 1430
    r = (x1 - x0) // 2
    d.arc((x0, y0, x1, y0 + 2 * r), 180, 360, fill=col, width=sw)
    d.line((x0 + sw // 2, y0 + r, x0 + sw // 2, y1), fill=col, width=sw)
    d.line((x1 - sw // 2, y0 + r, x1 - sw // 2, y1), fill=col, width=sw)
    d.line((x0, y1 - sw // 2, x1, y1 - sw // 2), fill=col, width=sw)
    d.line((640, 760, 900, 500), fill=col, width=sw)
    d.line((660, 960, 980, 640), fill=col, width=sw)
    im.resize((400, 400), Image.LANCZOS).save(BUILD / "icon_espelho.png")

    im = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.rectangle((200, 560, 1400, 600), fill=col)
    d.rectangle((150, 620, 1450, 680), fill=col)
    for k in range(8):
        cx = 230 + k * 163
        d.rectangle((cx - 22, 470, cx + 22, 556), fill=col)
    d.rectangle((330, 680, 380, 1350), fill=col)
    d.rectangle((1220, 680, 1270, 1350), fill=col)
    d.rectangle((150, 1350, 1450, 1384), fill=col)
    im.resize((400, 400), Image.LANCZOS).save(BUILD / "icon_pergolado.png")


def make_qr():
    import segno
    q = segno.make(WA, error="h")
    q.save(str(BUILD / "qr_whatsapp.png"), scale=24, border=0, dark=INK, light="#FFFFFF")


def grad_png(e):
    key = hashlib.md5(repr((e["a0"], e["a1"], e["color"], e["ease"])).encode()).hexdigest()[:10]
    path = BUILD / f"grad_{key}.png"
    if not path.exists():
        hgt = 1024
        r, g, b = hexrgb(e["color"])
        im = Image.new("RGBA", (8, hgt))
        px = im.load()
        for yy in range(hgt):
            t = (yy / (hgt - 1)) ** e["ease"] if e["a1"] > e["a0"] else \
                1 - (1 - yy / (hgt - 1)) ** e["ease"]
            a = e["a0"] + (e["a1"] - e["a0"]) * t
            for xx in range(8):
                px[xx, yy] = (r, g, b, round(255 * a))
        im.save(path)
    return path


def crop_img(e, dpi, quality, outdir):
    src = Image.open(ROOT / f"fotos/{e['src']}.jpg").convert("RGB")
    ar = e["w"] / e["h"]
    sw, sh = src.size
    if sw / sh > ar:
        cw, ch = sh * ar, sh
    else:
        cw, ch = sw, sw / ar
    fx, fy = e["focus"]
    left = min(max(fx * sw - cw / 2, 0), sw - cw)
    top = min(max(fy * sh - ch / 2, 0), sh - ch)
    im = src.crop((round(left), round(top), round(left + cw), round(top + ch)))
    tw = round(e["w"] / 25.4 * dpi)
    if tw < im.width:
        im = im.resize((tw, round(tw / ar)), Image.LANCZOS)
    im = im.filter(ImageFilter.UnsharpMask(radius=1.2, percent=35, threshold=2))
    name = f"{e['src']}_{round(e['w'])}x{round(e['h'])}_{hashlib.md5(repr(e['focus']).encode()).hexdigest()[:4]}.jpg"
    path = outdir / name
    im.save(path, quality=quality, optimize=True, progressive=True)
    return path


# ---------------------------------------------------------------- verificação de texto

def runs_of(paragraph, color, accent):
    out = []
    for part in re.split(r"(\{[^}]*\})", paragraph):
        if not part:
            continue
        if part.startswith("{"):
            out.append((part[1:-1], accent))
        else:
            out.append((part, color))
    return out


def check_text(e, pageno):
    f = ImageFont.truetype(str(FONTS[(e["font"], e["bold"])]), 1000)
    scale = e["size"] / 1000
    maxw = e["w"] / 25.4 * 72
    lines = 0
    widest = 0
    for para in e["text"].split("\n"):
        plain = para.replace("{", "").replace("}", "")
        words = plain.split(" ")
        cur = ""
        for wd in words:
            cand = (cur + " " + wd) if cur else wd
            wpt = f.getlength(cand) * scale + len(cand) * e["track"] * e["size"]
            if wpt > maxw and cur:
                lines += 1
                widest = max(widest, f.getlength(cur) * scale)
                cur = wd
            else:
                cur = cand
        lines += 1
        widest = max(widest, f.getlength(cur) * scale + len(cur) * e["track"] * e["size"])
    need = lines * e["size"] * e["lh"] + (len(e["text"].split("\n")) - 1) * e["pspace"]
    have = e["h"] / 25.4 * 72
    warn = []
    if need > have + 0.5:
        warn.append(f"p{pageno} altura: precisa {need / 72 * 25.4:.1f}mm, tem {e['h']:.1f}mm")
    if widest > maxw + 0.5:
        warn.append(f"p{pageno} largura: {widest / 72 * 25.4:.1f}mm > {e['w']:.1f}mm")
    for w_ in warn:
        print("  AVISO", w_, "->", e["text"][:50].replace("\n", " / "))
    return not warn


# ---------------------------------------------------------------- HTML / PDF

def mm(v):
    return f"{v:.3f}mm"


def html_el(e, assets, outdir):
    k = e["k"]
    box = f"left:{mm(e['x'])};top:{mm(e['y'])};width:{mm(e['w'])};height:{mm(e['h'])};"
    if k == "rect":
        r, g, b = hexrgb(e["fill"])
        st = f"background:rgba({r},{g},{b},{e['alpha']});border-radius:{mm(e['radius'])};"
        if e["stroke"]:
            st += f"border:{mm(e['sw'])} solid {e['stroke']};"
        out = f'<div class="e" style="{box}{st}"></div>'
    elif k == "oval":
        st = "border-radius:50%;"
        if e["fill"]:
            st += f"background:{e['fill']};"
        if e["stroke"]:
            st += f"border:{mm(e['sw'])} solid {e['stroke']};"
        out = f'<div class="e" style="{box}{st}"></div>'
    elif k in ("img", "png", "grad"):
        path = assets(e)
        rel = Path(shutil.copy(path, outdir / path.name)).name
        out = f'<img class="e" style="{box}" src="{rel}" alt="">'
    elif k == "text":
        jc = {"top": "flex-start", "middle": "center", "bottom": "flex-end"}[e["valign"]]
        st = (f"font-family:'{FAMILY[e['font']]}';font-weight:{700 if e['bold'] else 400};"
              f"font-size:{e['size']}pt;line-height:{e['lh']};color:{e['color']};"
              f"text-align:{e['align']};letter-spacing:{e['track']}em;justify-content:{jc};")
        paras = []
        for i, para in enumerate(e["text"].split("\n")):
            inner = "".join(
                html.escape(t) if c == e["color"] else
                f'<span style="color:{c}">{html.escape(t)}</span>'
                for t, c in runs_of(para, e["color"], e["accent"]))
            mt = f' style="margin-top:{e["pspace"]}pt"' if i and e["pspace"] else ""
            paras.append(f"<p{mt}>{inner or '&nbsp;'}</p>")
        out = f'<div class="e t" style="{box}{st}">{"".join(paras)}</div>'
    elif k == "link":
        out = ""
    if e.get("link"):
        out += (f'<a class="e" href="{html.escape(e["link"])}" style="{box}" '
                f'aria-label="WhatsApp {CONTATO_NOME}"></a>')
    return out


def build_html(pages, dpi, quality, tag):
    outdir = BUILD / tag
    if outdir.exists():
        shutil.rmtree(outdir)
    outdir.mkdir(parents=True)
    for f in FONTS.values():
        shutil.copy(f, outdir / f.name)

    def assets(e):
        if e["k"] == "img":
            return crop_img(e, dpi, quality, BUILD / f"img_{tag}")
        if e["k"] == "grad":
            return grad_png(e)
        return Path(e["path"])

    (BUILD / f"img_{tag}").mkdir(exist_ok=True)
    faces = "".join(
        f"@font-face{{font-family:'{FAMILY[fam]}';font-weight:{700 if b else 400};"
        f"src:url('{p.name}')}}" for (fam, b), p in FONTS.items())
    body = []
    for pg in pages:
        els = "".join(html_el(e, assets, outdir) for e in pg["els"])
        body.append(f'<section class="page" style="background:{pg["bg"]}">{els}</section>')
    doc = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">
<title>JP Alumínios | Apresentação institucional</title>
<style>{faces}
@page{{size:210mm 297mm;margin:0}}
html,body{{margin:0;padding:0;background:{INK}}}
*{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.page{{position:relative;width:210mm;height:297mm;overflow:hidden;break-after:page}}
.page:last-child{{break-after:auto}}
.e{{position:absolute;box-sizing:border-box;margin:0;display:block}}
.t{{display:flex;flex-direction:column}}
.t p{{margin:0}}
</style></head><body>{''.join(body)}</body></html>"""
    path = outdir / "index.html"
    path.write_text(doc, encoding="utf-8")
    return path


def print_pdf(html_path, pdf_path):
    subprocess.run([CHROME, "--headless=new", "--no-sandbox", "--disable-gpu",
                    "--allow-file-access-from-files", "--no-pdf-header-footer",
                    "--virtual-time-budget=10000", f"--print-to-pdf={pdf_path}",
                    html_path.as_uri()], check=True, capture_output=True)


def set_pdf_meta(pdf_path):
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(str(pdf_path))
    w = PdfWriter(clone_from=r)
    w.add_metadata({"/Title": "JP Alumínios | Apresentação institucional",
                    "/Author": "JP Alumínios", "/Subject": "Vidros e esquadrias de alumínio",
                    "/Keywords": f"orçamento, {CONTATO_NOME}, {CONTATO_TEL}"})
    w.compress_identical_objects()
    with open(pdf_path, "wb") as fh:
        w.write(fh)


# ---------------------------------------------------------------- PPTX

def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def pptx_fill(shape, color, alpha=1.0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*hexrgb(color))
    if alpha < 1:
        clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = clr.makeelement(qn("a:alpha"), {"val": str(round(alpha * 100000))})
        clr.append(a)


def build_pptx(pages, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(W), Mm(H)
    prs.core_properties.title = "JP Alumínios | Apresentação institucional"
    prs.core_properties.author = "JP Alumínios"
    layout = prs.slide_layouts[6]

    def assets(e):
        if e["k"] == "img":
            return crop_img(e, 200, 88, BUILD / "img_pptx")

    (BUILD / "img_pptx").mkdir(exist_ok=True)
    for pg in pages:
        s = prs.slides.add_slide(layout)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(*hexrgb(pg["bg"]))
        for e in pg["els"]:
            x, y, w, h = Mm(e["x"]), Mm(e["y"]), Mm(e["w"]), Mm(e["h"])
            k = e["k"]
            shp = None
            if k == "rect":
                kind = MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE
                shp = s.shapes.add_shape(kind, x, y, w, h)
                strip_style(shp)
                if e["radius"]:
                    shp.adjustments[0] = min(0.5, e["radius"] / min(e["w"], e["h"]))
                pptx_fill(shp, e["fill"], e["alpha"])
                if e["stroke"]:
                    shp.line.color.rgb = RGBColor(*hexrgb(e["stroke"]))
                    shp.line.width = Mm(e["sw"])
                else:
                    shp.line.fill.background()
            elif k == "oval":
                shp = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
                strip_style(shp)
                if e["fill"]:
                    pptx_fill(shp, e["fill"])
                else:
                    shp.fill.background()
                if e["stroke"]:
                    shp.line.color.rgb = RGBColor(*hexrgb(e["stroke"]))
                    shp.line.width = Mm(e["sw"])
                else:
                    shp.line.fill.background()
            elif k == "img":
                shp = s.shapes.add_picture(str(assets(e)), x, y, w, h)
            elif k == "png":
                shp = s.shapes.add_picture(e["path"], x, y, w, h)
            elif k == "grad":
                shp = s.shapes.add_picture(str(grad_png(e)), x, y, w, h)
            elif k == "text":
                shp = s.shapes.add_textbox(x, y, w, h)
                tf = shp.text_frame
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                                      "bottom": MSO_ANCHOR.BOTTOM}[e["valign"]]
                for i, para in enumerate(e["text"].split("\n")):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                                   "right": PP_ALIGN.RIGHT}[e["align"]]
                    p.line_spacing = Pt(e["size"] * e["lh"])
                    if i and e["pspace"]:
                        p.space_before = Pt(e["pspace"])
                    for t, c in runs_of(para, e["color"], e["accent"]):
                        r = p.add_run()
                        r.text = t
                        f = r.font
                        f.name = FAMILY[e["font"]]
                        f.size = Pt(e["size"])
                        f.bold = e["bold"]
                        f.color.rgb = RGBColor(*hexrgb(c))
                        if e["track"]:
                            r._r.get_or_add_rPr().set("spc", str(round(e["track"] * e["size"] * 100)))
            elif k == "link":
                shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
                strip_style(shp)
                shp.fill.background()
                shp.line.fill.background()
            if shp is not None and e.get("link"):
                shp.click_action.hyperlink.address = e["link"]
    prs.save(path)


# ---------------------------------------------------------------- principal

def main():
    BUILD.mkdir(exist_ok=True)
    OUT.mkdir(exist_ok=True)
    make_icons()
    make_qr()
    pages = [fn() for fn in PAGES]
    ok = all([check_text(e, i + 1) for i, pg in enumerate(pages) for e in pg["els"]
              if e["k"] == "text"])
    print("textos cabem nas caixas" if ok else "ATENÇÃO: textos com risco de corte")

    hq = OUT / f"{ARQ}.pdf"
    lite = OUT / f"{ARQ}-leve.pdf"
    print_pdf(build_html(pages, 230, 90, "hq"), hq)
    print_pdf(build_html(pages, 120, 72, "lite"), lite)
    for p in (hq, lite):
        set_pdf_meta(p)
    build_pptx(pages, OUT / f"{ARQ}-editavel.pptx")

    jpg_dir = OUT / "paginas-para-whatsapp"
    if jpg_dir.exists():
        shutil.rmtree(jpg_dir)
    jpg_dir.mkdir()
    subprocess.run(["pdftoppm", "-jpeg", "-jpegopt", "quality=88,progressive=y",
                    "-scale-to-x", "1240", "-scale-to-y", "-1", str(hq),
                    str(jpg_dir / "pagina")], check=True)
    for f in OUT.rglob("*"):
        if f.is_file():
            print(f"  {f.relative_to(OUT)}  {f.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main()
