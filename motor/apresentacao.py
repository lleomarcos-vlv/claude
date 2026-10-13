"""Motor de apresentações: uma definição de páginas gera PDF, PDF leve, PPTX editável e JPGs."""
import hashlib
import html
import re
import shutil
import subprocess
from pathlib import Path

from PIL import Image, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

CHROME = "/opt/pw-browsers/chromium"
W, H = 210, 297


# ---------------------------------------------------------------- elementos (medidas em mm)

def R(x, y, w, h, fill, alpha=1.0, radius=0, stroke=None, sw=0.3, link=None):
    return dict(k="rect", x=x, y=y, w=w, h=h, fill=fill, alpha=alpha, radius=radius,
                stroke=stroke, sw=sw, link=link)


def O(x, y, d, fill=None, stroke=None, sw=0.4):
    return dict(k="oval", x=x, y=y, w=d, h=d, fill=fill, stroke=stroke, sw=sw)


def I(x, y, w, h, src, focus=(0.5, 0.5), link=None):
    """Foto recortada para preencher a caixa (object-fit: cover)."""
    return dict(k="img", x=x, y=y, w=w, h=h, src=str(src), focus=focus, link=link)


def P(x, y, w, h, path, link=None, keep=False, bg=None):
    """Imagem esticada na caixa; keep=True preserva o arquivo (ex.: QR Code); bg achata a
    transparência sobre essa cor para gerar JPEG leve."""
    return dict(k="png", x=x, y=y, w=w, h=h, path=str(path), link=link, keep=keep, bg=bg)


def fit(path, x, y, w, h, align="center", valign="bottom", link=None, bg=None):
    """Imagem inteira, proporcional, dentro da caixa (object-fit: contain)."""
    with Image.open(path) as im:
        ar = im.width / im.height
    iw, ih = (w, w / ar) if w / h < ar else (h * ar, h)
    ix = x + {"left": 0, "center": (w - iw) / 2, "right": w - iw}[align]
    iy = y + {"top": 0, "middle": (h - ih) / 2, "bottom": h - ih}[valign]
    return P(ix, iy, iw, ih, path, link=link, bg=bg)


def G(x, y, w, h, a0, a1, color="#000000", ease=1.6):
    """Degradê vertical de uma cor, da opacidade a0 (topo) para a1 (base)."""
    return dict(k="grad", x=x, y=y, w=w, h=h, a0=a0, a1=a1, color=color, ease=ease)


def T(x, y, w, h, text, size, color="#000000", font="body", bold=False, italic=False,
      align="left", valign="top", lh=1.35, track=0.0, accent="#000000", pspace=0, link=None):
    """Texto; \\n separa parágrafos e {trecho} usa a cor de destaque."""
    return dict(k="text", x=x, y=y, w=w, h=h, text=text, size=size, color=color, font=font,
                bold=bold, italic=italic, align=align, valign=valign, lh=lh, track=track,
                accent=accent, pspace=pspace, link=link)


def A(x, y, w, h, url):
    return dict(k="link", x=x, y=y, w=w, h=h, link=url)


def hexrgb(c):
    c = c.lstrip("#")
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


def qr_png(url, path, dark="#000000"):
    import segno
    segno.make(url, error="h").save(str(path), scale=24, border=0, dark=dark, light="#FFFFFF")
    return path


def runs_of(paragraph, color, accent):
    out = []
    for part in re.split(r"(\{[^}]*\})", paragraph):
        if part:
            out.append((part[1:-1], accent) if part.startswith("{") else (part, color))
    return out


def mm(v):
    return f"{v:.3f}mm"


# ---------------------------------------------------------------- projeto

class Projeto:
    def __init__(self, root, out, arquivo, titulo, autor, assunto, palavras, fontes, familias,
                 fundo, rotulo_link):
        self.root = Path(root)
        self.out = Path(out)
        self.build = self.root / "build"
        self.arquivo = arquivo
        self.titulo = titulo
        self.autor = autor
        self.assunto = assunto
        self.palavras = palavras
        self.fontes = fontes
        self.familias = familias
        self.fundo = fundo
        self.rotulo_link = rotulo_link

    # ---- imagens

    def grad_png(self, e):
        key = hashlib.md5(repr((e["a0"], e["a1"], e["color"], e["ease"])).encode()).hexdigest()[:10]
        path = self.build / f"grad_{key}.png"
        if not path.exists():
            hgt = 1024
            r, g, b = hexrgb(e["color"])
            im = Image.new("RGBA", (8, hgt))
            px = im.load()
            for yy in range(hgt):
                t = (yy / (hgt - 1)) ** e["ease"] if e["a1"] > e["a0"] else \
                    1 - (1 - yy / (hgt - 1)) ** e["ease"]
                a = e["a0"] + (e["a1"] - e["a0"]) * t
                for xx in range(8):
                    px[xx, yy] = (r, g, b, round(255 * a))
            im.save(path)
        return path

    def src_path(self, src):
        p = Path(src)
        return p if p.suffix else self.root / f"fotos/{src}.jpg"

    def crop_img(self, e, dpi, quality, outdir):
        src = Image.open(self.src_path(e["src"])).convert("RGB")
        ar = e["w"] / e["h"]
        sw, sh = src.size
        cw, ch = (sh * ar, sh) if sw / sh > ar else (sw, sw / ar)
        fx, fy = e["focus"]
        left = min(max(fx * sw - cw / 2, 0), sw - cw)
        top = min(max(fy * sh - ch / 2, 0), sh - ch)
        im = src.crop((round(left), round(top), round(left + cw), round(top + ch)))
        tw = round(e["w"] / 25.4 * dpi)
        if tw < im.width:
            im = im.resize((tw, round(tw / ar)), Image.LANCZOS)
        im = im.filter(ImageFilter.UnsharpMask(radius=1.2, percent=35, threshold=2))
        tag = hashlib.md5(repr((e["src"], e["focus"])).encode()).hexdigest()[:6]
        path = outdir / f"{Path(e['src']).stem}_{round(e['w'])}x{round(e['h'])}_{tag}.jpg"
        im.save(path, quality=quality, optimize=True, progressive=True)
        return path

    def place_img(self, e, dpi, quality, outdir):
        src = Path(e["path"])
        if e.get("keep"):
            return src
        im = Image.open(src)
        need = round(e["w"] / 25.4 * dpi * 1.5)
        alpha = (im.mode in ("RGBA", "LA") or "transparency" in im.info) and \
            im.convert("RGBA").getchannel("A").getextrema()[0] < 255
        jpg = src.suffix.lower() in (".jpg", ".jpeg")
        if im.width <= need and not e.get("bg") and (alpha or jpg):
            return src
        if im.width > need:
            im = im.resize((need, max(1, round(im.height * need / im.width))), Image.LANCZOS)
        tag = hashlib.md5(repr((str(src), need, quality, e.get("bg"))).encode()).hexdigest()[:6]
        if alpha and e.get("bg"):
            flat = Image.new("RGB", im.size, hexrgb(e["bg"]))
            flat.paste(im.convert("RGBA"), mask=im.convert("RGBA").getchannel("A"))
            im, alpha = flat, False
        if alpha:
            path = outdir / f"{src.stem}_{tag}.png"
            im.convert("RGBA").save(path, optimize=True)
        else:
            path = outdir / f"{src.stem}_{tag}.jpg"
            im.convert("RGB").save(path, quality=quality, optimize=True, progressive=True)
        return path

    # ---- verificação

    def check_text(self, e, pageno):
        f = ImageFont.truetype(str(self.fontes[(e["font"], e["bold"], e["italic"])]), 1000)
        scale = e["size"] / 1000
        maxw = e["w"] / 25.4 * 72
        lines = 0
        widest = 0
        for para in e["text"].split("\n"):
            plain = para.replace("{", "").replace("}", "")
            cur = ""
            for wd in plain.split(" "):
                cand = (cur + " " + wd) if cur else wd
                wpt = f.getlength(cand) * scale + len(cand) * e["track"] * e["size"]
                if wpt > maxw and cur:
                    lines += 1
                    widest = max(widest, f.getlength(cur) * scale)
                    cur = wd
                else:
                    cur = cand
            lines += 1
            widest = max(widest, f.getlength(cur) * scale + len(cur) * e["track"] * e["size"])
        need = lines * e["size"] * e["lh"] + (len(e["text"].split("\n")) - 1) * e["pspace"]
        have = e["h"] / 25.4 * 72
        warn = []
        if need > have + 0.5:
            warn.append(f"p{pageno} altura: precisa {need / 72 * 25.4:.1f}mm, tem {e['h']:.1f}mm")
        if widest > maxw + 0.5:
            warn.append(f"p{pageno} largura: {widest / 72 * 25.4:.1f}mm > {e['w']:.1f}mm")
        for w_ in warn:
            print("  AVISO", w_, "->", e["text"][:50].replace("\n", " / "))
        return not warn

    # ---- HTML / PDF

    def html_el(self, e, assets, outdir):
        k = e["k"]
        box = f"left:{mm(e['x'])};top:{mm(e['y'])};width:{mm(e['w'])};height:{mm(e['h'])};"
        out = ""
        if k == "rect":
            r, g, b = hexrgb(e["fill"])
            st = f"background:rgba({r},{g},{b},{e['alpha']});border-radius:{mm(e['radius'])};"
            if e["stroke"]:
                st += f"border:{mm(e['sw'])} solid {e['stroke']};"
            out = f'<div class="e" style="{box}{st}"></div>'
        elif k == "oval":
            st = "border-radius:50%;"
            if e["fill"]:
                st += f"background:{e['fill']};"
            if e["stroke"]:
                st += f"border:{mm(e['sw'])} solid {e['stroke']};"
            out = f'<div class="e" style="{box}{st}"></div>'
        elif k in ("img", "png", "grad"):
            path = assets(e)
            rel = Path(shutil.copy(path, outdir / path.name)).name
            out = f'<img class="e" style="{box}" src="{rel}" alt="">'
        elif k == "text":
            jc = {"top": "flex-start", "middle": "center", "bottom": "flex-end"}[e["valign"]]
            st = (f"font-family:'{self.familias[e['font']]}';"
                  f"font-weight:{700 if e['bold'] else 400};"
                  f"font-style:{'italic' if e['italic'] else 'normal'};"
                  f"font-size:{e['size']}pt;line-height:{e['lh']};color:{e['color']};"
                  f"text-align:{e['align']};letter-spacing:{e['track']}em;justify-content:{jc};")
            paras = []
            for i, para in enumerate(e["text"].split("\n")):
                inner = "".join(
                    html.escape(t) if c == e["color"] else
                    f'<span style="color:{c}">{html.escape(t)}</span>'
                    for t, c in runs_of(para, e["color"], e["accent"]))
                mt = f' style="margin-top:{e["pspace"]}pt"' if i and e["pspace"] else ""
                paras.append(f"<p{mt}>{inner or '&nbsp;'}</p>")
            out = f'<div class="e t" style="{box}{st}">{"".join(paras)}</div>'
        if e.get("link"):
            out += (f'<a class="e" href="{html.escape(e["link"])}" style="{box}" '
                    f'aria-label="{html.escape(self.rotulo_link)}"></a>')
        return out

    def build_html(self, pages, dpi, quality, tag):
        outdir = self.build / tag
        if outdir.exists():
            shutil.rmtree(outdir)
        outdir.mkdir(parents=True)
        imgdir = self.build / f"img_{tag}"
        imgdir.mkdir(exist_ok=True)
        for f in self.fontes.values():
            shutil.copy(f, outdir / f.name)

        def assets(e):
            if e["k"] == "img":
                return self.crop_img(e, dpi, quality, imgdir)
            if e["k"] == "grad":
                return self.grad_png(e)
            return self.place_img(e, dpi, quality, imgdir)

        faces = "".join(
            f"@font-face{{font-family:'{self.familias[fam]}';font-weight:{700 if b else 400};"
            f"font-style:{'italic' if it else 'normal'};src:url('{p.name}')}}"
            for (fam, b, it), p in self.fontes.items())
        body = []
        for pg in pages:
            els = "".join(self.html_el(e, assets, outdir) for e in pg["els"])
            body.append(f'<section class="page" style="background:{pg["bg"]}">{els}</section>')
        doc = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">
<title>{html.escape(self.titulo)}</title>
<style>{faces}
@page{{size:{W}mm {H}mm;margin:0}}
html,body{{margin:0;padding:0;background:{self.fundo}}}
*{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.page{{position:relative;width:{W}mm;height:{H}mm;overflow:hidden;break-after:page}}
.page:last-child{{break-after:auto}}
.e{{position:absolute;box-sizing:border-box;margin:0;display:block}}
.t{{display:flex;flex-direction:column}}
.t p{{margin:0}}
</style></head><body>{''.join(body)}</body></html>"""
        path = outdir / "index.html"
        path.write_text(doc, encoding="utf-8")
        return path

    @staticmethod
    def print_pdf(html_path, pdf_path):
        subprocess.run([CHROME, "--headless=new", "--no-sandbox", "--disable-gpu",
                        "--allow-file-access-from-files", "--no-pdf-header-footer",
                        "--virtual-time-budget=10000", f"--print-to-pdf={pdf_path}",
                        html_path.as_uri()], check=True, capture_output=True)

    def set_pdf_meta(self, pdf_path):
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(clone_from=PdfReader(str(pdf_path)))
        w.add_metadata({"/Title": self.titulo, "/Author": self.autor, "/Subject": self.assunto,
                        "/Keywords": self.palavras})
        w.compress_identical_objects()
        with open(pdf_path, "wb") as fh:
            w.write(fh)

    # ---- PPTX

    @staticmethod
    def _strip_style(shape):
        st = shape._element.find(qn("p:style"))
        if st is not None:
            shape._element.remove(st)

    @staticmethod
    def _fill(shape, color, alpha=1.0):
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*hexrgb(color))
        if alpha < 1:
            clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            clr.append(clr.makeelement(qn("a:alpha"), {"val": str(round(alpha * 100000))}))

    def _line(self, shp, e):
        if e["stroke"]:
            shp.line.color.rgb = RGBColor(*hexrgb(e["stroke"]))
            shp.line.width = Mm(e["sw"])
        else:
            shp.line.fill.background()

    def build_pptx(self, pages, path):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Mm(W), Mm(H)
        prs.core_properties.title = self.titulo
        prs.core_properties.author = self.autor
        layout = prs.slide_layouts[6]
        imgdir = self.build / "img_pptx"
        imgdir.mkdir(exist_ok=True)
        for pg in pages:
            s = prs.slides.add_slide(layout)
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = RGBColor(*hexrgb(pg["bg"]))
            for e in pg["els"]:
                x, y, w, h = Mm(e["x"]), Mm(e["y"]), Mm(e["w"]), Mm(e["h"])
                k = e["k"]
                shp = None
                if k == "rect":
                    kind = MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE
                    shp = s.shapes.add_shape(kind, x, y, w, h)
                    self._strip_style(shp)
                    if e["radius"]:
                        shp.adjustments[0] = min(0.5, e["radius"] / min(e["w"], e["h"]))
                    self._fill(shp, e["fill"], e["alpha"])
                    self._line(shp, e)
                elif k == "oval":
                    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
                    self._strip_style(shp)
                    if e["fill"]:
                        self._fill(shp, e["fill"])
                    else:
                        shp.fill.background()
                    self._line(shp, e)
                elif k == "img":
                    shp = s.shapes.add_picture(str(self.crop_img(e, 200, 88, imgdir)), x, y, w, h)
                elif k == "png":
                    shp = s.shapes.add_picture(str(self.place_img(e, 200, 88, imgdir)), x, y, w, h)
                elif k == "grad":
                    shp = s.shapes.add_picture(str(self.grad_png(e)), x, y, w, h)
                elif k == "text":
                    shp = s.shapes.add_textbox(x, y, w, h)
                    tf = shp.text_frame
                    tf.word_wrap = True
                    tf.auto_size = MSO_AUTO_SIZE.NONE
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                                          "bottom": MSO_ANCHOR.BOTTOM}[e["valign"]]
                    for i, para in enumerate(e["text"].split("\n")):
                        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                                       "right": PP_ALIGN.RIGHT}[e["align"]]
                        p.line_spacing = Pt(e["size"] * e["lh"])
                        if i and e["pspace"]:
                            p.space_before = Pt(e["pspace"])
                        for t, c in runs_of(para, e["color"], e["accent"]):
                            r = p.add_run()
                            r.text = t
                            f = r.font
                            f.name = self.familias[e["font"]]
                            f.size = Pt(e["size"])
                            f.bold = e["bold"]
                            f.italic = e["italic"]
                            f.color.rgb = RGBColor(*hexrgb(c))
                            if e["track"]:
                                r._r.get_or_add_rPr().set(
                                    "spc", str(round(e["track"] * e["size"] * 100)))
                elif k == "link":
                    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
                    self._strip_style(shp)
                    shp.fill.background()
                    shp.line.fill.background()
                if shp is not None and e.get("link"):
                    shp.click_action.hyperlink.address = e["link"]
        prs.save(path)

    # ---- tudo

    def gerar(self, pages):
        self.build.mkdir(exist_ok=True)
        self.out.mkdir(parents=True, exist_ok=True)
        ok = all([self.check_text(e, i + 1) for i, pg in enumerate(pages) for e in pg["els"]
                  if e["k"] == "text"])
        print("textos cabem nas caixas" if ok else "ATENÇÃO: textos com risco de corte")
        hq = self.out / f"{self.arquivo}.pdf"
        lite = self.out / f"{self.arquivo}-leve.pdf"
        self.print_pdf(self.build_html(pages, 230, 90, "hq"), hq)
        self.print_pdf(self.build_html(pages, 120, 72, "lite"), lite)
        for p in (hq, lite):
            self.set_pdf_meta(p)
        self.build_pptx(pages, self.out / f"{self.arquivo}-editavel.pptx")
        jpg_dir = self.out / "paginas-para-whatsapp"
        if jpg_dir.exists():
            shutil.rmtree(jpg_dir)
        jpg_dir.mkdir()
        subprocess.run(["pdftoppm", "-jpeg", "-jpegopt", "quality=88,progressive=y",
                        "-scale-to-x", "1240", "-scale-to-y", "-1", str(hq),
                        str(jpg_dir / "pagina")], check=True)
        for f in sorted(self.out.rglob("*")):
            if f.is_file():
                print(f"  {f.relative_to(self.out)}  {f.stat().st_size / 1024:.0f} KB")
